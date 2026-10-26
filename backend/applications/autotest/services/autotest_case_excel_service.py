# -*- coding: utf-8 -*-
"""
@Author  : yangkai
@Email   : 807440781@qq.com
@Project : KeenRunner
@Module  : autotest_case_excel_service.py
@DateTime: 2026/8/1
"""
from __future__ import annotations

import asyncio
import io
import os
import re
import traceback
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple
from xml.etree import ElementTree
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, PatternFill, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet
from tortoise.transactions import in_transaction

from backend.applications.autotest.models.autotest_case_model import AutoTestCaseModel
from backend.applications.autotest.models.autotest_project_model import AutoTestProjectModel
from backend.applications.autotest.models.autotest_step_model import AutoTestStepModel
from backend.applications.autotest.services.autotest_case_excel_codec import CaseScriptCodec
from backend.configure import LOGGER, PROJECT_CONFIG
from backend.enums import (
    AutoTestCaseAttr,
    AutoTestCaseType,
    AutoTestReqArgsType,
    AutoTestStepType,
    AutoTestTcpDatagramLength,
    AutoTestTcpEndSymbol,
    AutoTestTcpRadix,
    HTTPMethod,
)
from backend.services import get_current_username

__all__ = ["CaseExcelService"]


class CaseExcelService:
    """
    公共接口 Excel 业务服务。

    职责分组：
    - 通用样式与文件工具：列宽自适应、数据区样式、导出文件名、用例批量装载
    - 通道A 报文矩阵导出：HEAD/BODY JSONPath 矩阵、一个用例对应一个Sheet页、目录页内跳转
    - 数据源sheet样式：分区标记黄底、前导'转 quotePrefix（供数据源/性能模块视图复用）
    - 通道B 脚本导出：行序列化与往返安全检测、按协议组合选模板并分sheet写入
    - 通道C 脚本导入：首sheet布局识别、行解析校验、批量匹配后单事务落库
    """

    # ==================== 常量：协议标识与导出样式 ====================
    TCP_PROTOCOL: str = "TCP"
    HTTP_PROTOCOL: str = "HTTP"
    GENERATE_DIRECTORY_THRESHOLD: int = 2
    SHEET_TITLE_MAX_LENGTH: int = 31
    SHEET_STYLE_ROW_HEIGHT: int = 40
    SHEET_STYLE_COL_WIDTH_MIN: int = 8
    SHEET_STYLE_COL_WIDTH_MAX: int = 60
    SHEET_STYLE_FONT = Font(bold=True)
    SHEET_STYLE_SIDE = Side(style="thin", color="000000")
    SHEET_STYLE_PATTERN_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")
    SHEET_STYLE_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
    SHEET_STYLE_MARK_BORDER = Border(left=SHEET_STYLE_SIDE, right=SHEET_STYLE_SIDE, top=SHEET_STYLE_SIDE, bottom=SHEET_STYLE_SIDE)

    # 报文导出分区
    DATAGRAM_SECTIONS: Tuple[Tuple[str, str], ...] = (("HEAD", "head_pairs"), ("BODY", "body_pairs"))

    # 数据源分区标记
    DATAGRAM_SECTION_MARKS = frozenset({"HEAD", "BODY", "ASSERT_HEAD", "ASSERT_BODY"})

    # ==================== 常量：脚本模板布局 ====================
    TCP_SHEET_NAME = "TCP模板"
    TCP_SHEET_COLUMNS: Tuple[str, ...] = (
        "接口名称", "所属应用", "协议类型", "接口描述",
        "配置名称", "报文进制", "长度位", "报文总长度", "请求结束符", "响应结束符",
        "请求体类型", "请求体", "变量", "提取", "断言", "所属人员",
    )
    HTTP_SHEET_NAME = "HTTP模板"
    HTTP_SHEET_COLUMNS: Tuple[str, ...] = (
        "接口名称", "所属应用", "协议类型", "接口描述",
        "请求方式", "配置名称", "请求路径", "请求体类型", "请求体", "请求头",
        "变量", "提取", "断言", "所属人员",
    )
    SCRIPT_SHEET_LAYOUTS: Dict[str, Tuple[str, Tuple[str, ...]]] = {
        TCP_PROTOCOL: (TCP_SHEET_NAME, TCP_SHEET_COLUMNS),
        HTTP_PROTOCOL: (HTTP_SHEET_NAME, HTTP_SHEET_COLUMNS),
    }
    SCRIPT_EXPORT_SHEET_NAMES: Dict[str, str] = {
        TCP_PROTOCOL: "TCP数据",
        HTTP_PROTOCOL: "HTTP数据",
    }
    SCRIPT_DATA_START_ROW: int = 3
    TCP_FRAME_CONFIG_COLUMNS: Tuple[Tuple[str, str], ...] = (
        ("报文进制", "tcp_radix"),
        ("长度位", "tcp_radix_length"),
        ("报文总长度", "tcp_datagram_length"),
        ("请求结束符", "tcp_request_end_symbol"),
        ("响应结束符", "tcp_response_end_symbol"),
    )

    # ==================== 常量：TCP 帧配置校验 ====================
    TCP_RADIX_VALUES = frozenset(e.value for e in AutoTestTcpRadix)
    TCP_DATAGRAM_LENGTH_VALUES = frozenset(e.value for e in AutoTestTcpDatagramLength)
    TCP_REQUEST_END_SYMBOL_VALUES = frozenset(e.value for e in AutoTestTcpEndSymbol)
    TCP_RADIX_LENGTH_MIN, TCP_RADIX_LENGTH_MAX = 2, 10
    TCP_RESPONSE_END_SYMBOL_MAX = 64

    # ==================== 常量：脚本模板文件 ====================
    SCRIPT_TEMPLATE_FILES: Dict[Tuple[bool, bool], str] = {
        (True, False): "接口模板-HTTP.xlsx",
        (False, True): "接口模板-TCP.xlsx",
        (True, True): "接口模板-HTTP&TCP.xlsx",
    }
    SCRIPT_IMPORT_TEMPLATE_FILE = "接口模板-HTTP&TCP.xlsx"
    SCRIPT_IMPORT_TEMPLATE_DOWNLOAD_NAME = "接口模板.xlsx"

    # ==================== 常量：脚本导入校验集合 ====================
    HTTP_METHODS = frozenset(e.value for e in HTTPMethod)
    REQ_ARGS_TYPE = frozenset(e.value for e in AutoTestReqArgsType)
    TCP_ARGS_TYPE = frozenset({AutoTestReqArgsType.XML.value, AutoTestReqArgsType.JSON.value, AutoTestReqArgsType.RAW.value})

    # ==================== 通用样式与文件工具 ====================

    @staticmethod
    def _calculation_text_width(text: str) -> int:
        """计算文本占用单元格显示宽度：ASCII编码计1，中文计2。"""
        return sum(2 if ord(char) > 127 else 1 for char in text)

    @classmethod
    def _auto_size_sheet_columns(cls, sheet: Worksheet) -> None:
        """根据单元格内容自适应列宽，多行文本取最长行。"""
        max_widths: Dict[int, int] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                for line in str(cell.value).splitlines() or [""]:
                    width = cls._calculation_text_width(line)
                    if width > max_widths.get(cell.column, 0):
                        max_widths[cell.column] = width
        for col_idx in range(1, (sheet.max_column or 0) + 1):
            width = min(cls.SHEET_STYLE_COL_WIDTH_MAX, max(cls.SHEET_STYLE_COL_WIDTH_MIN, max_widths.get(col_idx, 0) + 2))
            sheet.column_dimensions[get_column_letter(col_idx)].width = width

    @classmethod
    def _set_sheet_cells_style(cls, sheet: Worksheet, *, start_row: int = 1, row_height: Optional[float] = SHEET_STYLE_ROW_HEIGHT) -> None:
        """数据区水平/垂直居中并设置四边边框。"""
        max_row = sheet.max_row or 0
        max_column = sheet.max_column or 0
        if max_row < start_row or max_column < 1:
            return
        for row in sheet.iter_rows(min_row=start_row, max_row=max_row, max_col=max_column):
            if row_height is not None:
                sheet.row_dimensions[row[0].row].height = row_height
            for cell in row:
                cell.alignment = cls.SHEET_STYLE_ALIGNMENT
                cell.border = cls.SHEET_STYLE_MARK_BORDER

    @staticmethod
    def _file_name(username: Optional[str], label: str) -> str:
        safe = re.sub(r'[\\/:*?"<>|\s]', "_", str(username or "").strip())
        prefix = f"{safe}_{label}" if safe else label
        stamp = datetime.now().strftime('%Y%m%d%H%M%S%f')[:-3]
        return f"{prefix}_{stamp}.xlsx"

    @staticmethod
    def _get(item: Any, name: str, default: Any = None) -> Any:
        """兼容dict/schema对象字段读取。"""
        return item.get(name, default) if isinstance(item, dict) else getattr(item, name, default)

    @staticmethod
    def _enum_val(raw: Any) -> str:
        """CharEnum/字符串 → 枚举值文本。"""
        if raw is None:
            return ""
        return getattr(raw, "value", raw) or ""

    # ==================== 公共接口装载（双通道共用） ====================

    @classmethod
    async def _load_public_api_cases(cls, case_ids: List[int], services: Any) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """
        两通道共用：公共接口 + 唯一步骤HTTP/TCP + 无数据源。

        :param case_ids: 用例主键列表
        :param services: CRUD聚合
        :return: (有效装载列表, 不合规明细[{case_id, case_name, reason}])
        """
        ordered_ids = list(dict.fromkeys(case_ids or []))
        if not ordered_ids:
            return [], []
        case_map: Dict[int, AutoTestCaseModel] = {
            case.id: case
            for case in await services.case_curd.model.filter(id__in=ordered_ids, state__not=1)
        }
        steps_by_case = await cls._load_case_steps(ordered_ids, services)
        project_names = await cls._load_project_names(
            [getattr(case, "case_project", None) for case in case_map.values()], services
        )

        valid: List[Dict[str, Any]] = []
        invalid: List[Dict[str, Any]] = []
        for case_id in ordered_ids:
            case = case_map.get(case_id)
            case_name = (getattr(case, "case_name", None) or str(case_id)) if case else str(case_id)
            if case is None:
                invalid.append({"case_id": case_id, "case_name": case_name, "reason": "用例不存在"})
                continue
            if getattr(case, "case_type", None) != AutoTestCaseType.PUBLIC_API:
                invalid.append({"case_id": case_id, "case_name": case_name, "reason": "用例类型不被允许"})
                continue
            own_steps = cls._collect_own_steps(steps_by_case.get(case_id, []))
            if len(own_steps) != 1:
                invalid.append({"case_id": case_id, "case_name": case_name, "reason": "用例步骤数异常"})
                continue
            step = own_steps[0]
            step_type = getattr(step, "step_type", None)
            if step_type not in (AutoTestStepType.HTTP, AutoTestStepType.TCP):
                invalid.append({"case_id": case_id, "case_name": case_name, "reason": "步骤类型不被允许"})
                continue
            if getattr(step, "data_source_id", None):
                invalid.append({"case_id": case_id, "case_name": case_name, "reason": "步骤存在数据源绑定"})
                continue
            project_id = getattr(case, "case_project", None)
            valid.append({
                "case_id": case_id,
                "case": case,
                "case_name": case_name,
                "step": step,
                "is_http": step_type == AutoTestStepType.HTTP,
                "project_name": project_names.get(project_id, ""),
            })
        return valid, invalid

    @staticmethod
    async def _load_case_steps(case_ids: List[int], services: Any) -> Dict[int, List[AutoTestStepModel]]:
        """查询全部启用步骤并按用例分组。"""
        steps: List[AutoTestStepModel] = await services.step_curd.model.filter(
            case_id__in=case_ids, state__not=1
        ).all()
        grouped: Dict[int, List[AutoTestStepModel]] = defaultdict(list)
        for step in steps:
            grouped[step.case_id].append(step)
        return grouped

    @staticmethod
    def _collect_own_steps(steps: List[AutoTestStepModel]) -> List[AutoTestStepModel]:
        """
        按 parent_step_id 自根向下收集用例自有步骤。

        与步骤树对齐：只统计从根可达的步骤，父步骤缺失的悬空子步骤不计入。
        该服务只消费公共接口的"唯一请求步骤"，无需构建完整树。
        """
        roots = [step for step in steps if getattr(step, "parent_step_id", None) is None]
        children: Dict[Any, List[AutoTestStepModel]] = defaultdict(list)
        for step in steps:
            parent_id = getattr(step, "parent_step_id", None)
            if parent_id is not None:
                children[parent_id].append(step)
        collected: List[AutoTestStepModel] = []
        pending = sorted(roots, key=lambda step: getattr(step, "step_no", 0))
        while pending:
            step = pending.pop()
            collected.append(step)
            pending.extend(children.get(getattr(step, "id", None), ()))
        return collected

    @staticmethod
    async def _load_project_names(project_ids: Iterable[Optional[int]], services: Any) -> Dict[Optional[int], str]:
        """批量查询应用名称；悬空引用回退空名（与逐条查询取不到时口径一致）。"""
        names: Dict[Optional[int], str] = {None: ""}
        ids = sorted({project_id for project_id in project_ids if project_id})
        if not ids:
            return names
        projects: List[AutoTestProjectModel] = await services.project_curd.model.filter(id__in=ids).all()
        names.update({project.id: project.project_name or "" for project in projects})
        for project_id in ids:
            names.setdefault(project_id, "")
        return names

    # ==================== 通道 A：导出报文数据（HEAD/BODY JSONPath 矩阵） ====================

    @classmethod
    def build_export_file_name(cls, username: Optional[str]) -> str:
        return cls._file_name(username, "接口报文数据")

    @classmethod
    def _kv_to_pairs(cls, kv_list: Optional[List[Any]]) -> List[Tuple[str, Any]]:
        pairs: List[Tuple[str, Any]] = []
        for item in kv_list or []:
            key = cls._get(item, "key")
            if key is None or not str(key).strip():
                continue
            pairs.append((f"$.{str(key).strip()}", cls._get(item, "value")))
        return pairs

    @classmethod
    def _flatten_jsonpath(cls, data: Any, prefix: str = "$") -> List[Tuple[str, Any]]:
        pairs: List[Tuple[str, Any]] = []
        if isinstance(data, dict):
            for key, value in data.items():
                path = f"{prefix}.{key}"
                if isinstance(value, (dict, list)) and value:
                    pairs.extend(cls._flatten_jsonpath(value, path))
                else:
                    pairs.append((path, "" if isinstance(value, (dict, list)) else value))
        elif isinstance(data, list):
            for index, value in enumerate(data):
                path = f"{prefix}[{index}]"
                if isinstance(value, (dict, list)) and value:
                    pairs.extend(cls._flatten_jsonpath(value, path))
                else:
                    pairs.append((path, "" if isinstance(value, (dict, list)) else value))
        return pairs

    @staticmethod
    def _xml_local_name(tag: str) -> str:
        """去掉Clark命名空间，仅保留本地标签名。"""
        if tag and "}" in tag:
            return tag.rsplit("}", 1)[-1]
        return tag or ""

    @classmethod
    def _flatten_xml_to_jsonpath_pairs(cls, xml_text: str) -> List[Tuple[str, Any]]:
        """
        按XML文档序将叶子字段展平为JSONPath列。

        路径约定与JSON展平对齐：$.Root.Child；同名兄弟仅在重复时加0-based下标；
        属性为$.Path.@attr。非法XML回退$.raw。
        """
        text = str(xml_text or "").strip()
        if not text:
            return []
        try:
            root = ElementTree.fromstring(text)
        except ElementTree.ParseError:
            return [("$.raw", xml_text)]

        pairs: List[Tuple[str, Any]] = []

        def walk(elem: ElementTree.Element, path: str) -> None:
            for attr_name, attr_value in elem.attrib.items():
                local_attr = cls._xml_local_name(attr_name)
                pairs.append((f"{path}.@{local_attr}", "" if attr_value is None else attr_value))

            children = list(elem)
            if not children:
                pairs.append((path, "" if elem.text is None else elem.text))
                return

            name_total = Counter(cls._xml_local_name(child.tag) for child in children)
            name_seen: Counter = Counter()
            for child in children:
                local = cls._xml_local_name(child.tag)
                name_seen[local] += 1
                if name_total[local] > 1:
                    child_path = f"{path}.{local}[{name_seen[local] - 1}]"
                else:
                    child_path = f"{path}.{local}"
                walk(child, child_path)

        root_name = cls._xml_local_name(root.tag) or "root"
        walk(root, f"$.{root_name}")
        return pairs

    @classmethod
    def _body_to_pairs(cls, step: Any) -> List[Tuple[str, Any]]:
        args = getattr(step, "request_args_type", None)
        if args == AutoTestReqArgsType.JSON:
            body = getattr(step, "request_body", None)
            return cls._flatten_jsonpath(body) if isinstance(body, dict) else []
        if args == AutoTestReqArgsType.XML:
            text = getattr(step, "request_text", None) or ""
            return cls._flatten_xml_to_jsonpath_pairs(text) if text else []
        if args == AutoTestReqArgsType.RAW:
            text = getattr(step, "request_text", None) or ""
            return [("$.raw", text)] if text else []
        attr = CaseScriptCodec.FORM_FIELD_BY_ARGS_TYPE.get(cls._enum_val(args))
        return cls._kv_to_pairs(getattr(step, attr, None)) if attr else []

    @classmethod
    async def prepare_export_cases(cls, case_ids: List[int], services: Any) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        loaded, invalid = await cls._load_public_api_cases(case_ids, services)
        valid = [{
            "case_name": item["case_name"],
            "protocol": cls.HTTP_PROTOCOL if item["is_http"] else cls.TCP_PROTOCOL,
            "project_name": item["project_name"],
            "case_desc": getattr(item["case"], "case_desc", None) or "",
            "created_user": getattr(item["case"], "created_user", None) or "",
            "head_pairs": cls._kv_to_pairs(getattr(item["step"], "request_header", None)),
            "body_pairs": cls._body_to_pairs(item["step"]),
        } for item in loaded]
        LOGGER.info(f"导出用例准备完成: 有效{len(valid)}个, 无效{len(invalid)}个")
        return valid, invalid

    @classmethod
    def build_export_workbook(cls, cases_data: List[Dict[str, Any]]) -> Workbook:
        """
        根据一个用例对应一个Sheet页构建导出工作簿。

        :param cases_data: 用例报文数据列表
        :return: 填充并完成样式的工作簿
        """
        workbook = Workbook()
        workbook.remove(workbook.active)
        used_titles: Set[str] = set()
        sheet_titles: List[Tuple[str, Dict[str, Any]]] = []
        for case_data in cases_data:
            title = cls._unique_sheet_title(case_data.get("case_name"), used_titles)
            used_titles.add(title)
            cls._fill_datagram_sheet(workbook.create_sheet(title=title), case_data)
            sheet_titles.append((title, case_data))

        # 如果Sheet页大于2时, 则创建目录Sheet页
        if len(sheet_titles) > cls.GENERATE_DIRECTORY_THRESHOLD:
            cls._build_datagram_directory(workbook, sheet_titles)

        for sheet in workbook.worksheets:
            cls._highlight_datagram_marks(sheet)
            cls._set_sheet_cells_style(sheet, start_row=1, row_height=cls.SHEET_STYLE_ROW_HEIGHT)
            cls._auto_size_sheet_columns(sheet)
        return workbook

    @classmethod
    def _unique_sheet_title(cls, case_name: Any, used_titles: Set[str]) -> str:
        """生成合法且未占用的sheet名：剔除Excel禁用字符，超长按 _序号 后缀去重。"""
        base = re.sub(r"[:\\/?*\[\]]", "_", str(case_name or "").strip()) or "用例"
        base = base[:cls.SHEET_TITLE_MAX_LENGTH]
        title, index = base, 1
        while title in used_titles:
            suffix = f"_{index}"
            title = base[:cls.SHEET_TITLE_MAX_LENGTH - len(suffix)] + suffix
            index += 1
        return title

    @classmethod
    def _fill_datagram_sheet(cls, sheet: Worksheet, case_data: Dict[str, Any]) -> None:
        """写入单用例报文：分区标记 + JSONPath 表头行与对应值行。"""
        header: List[Any] = []
        values: List[Any] = []
        for label, field in cls.DATAGRAM_SECTIONS:
            header.append(label)
            values.append("")
            for path, value in case_data.get(field) or []:
                header.append(path)
                values.append(
                    "" if value is None else (value if isinstance(value, (str, int, float, bool)) else str(value))
                )
        sheet.append(header)
        sheet.append(values)

    @classmethod
    def _build_datagram_directory(cls, workbook: Workbook, sheet_titles: List[Tuple[str, Dict[str, Any]]]) -> None:
        """构建目录sheet，并为接口名称列设置跨sheet内部跳转。"""
        directory = workbook.create_sheet(title="目录", index=0)
        directory.append(["序号", "接口名称", "所属应用", "接口描述", "所属人", "协议类型"])
        for index, (title, case_data) in enumerate(sheet_titles, start=1):
            directory.append([
                index, case_data.get("case_name"), case_data.get("project_name"),
                case_data.get("case_desc"), case_data.get("created_user"), case_data.get("protocol"),
            ])
            cell = directory.cell(row=index + 1, column=2)
            # 内部跳转链接必须用 location 属性（标准 OOXML 写法），Excel/WPS 均兼容；
            # 若写成 target="#'sheet'!A1"，openpyxl 会生成外部关系引用，WPS 无法识别跳转；
            # sheet 名中的单引号在引用中需双写转义
            escaped = title.replace("'", "''")
            cell.hyperlink = Hyperlink(
                ref=cell.coordinate,
                location=f"'{escaped}'!A1",
                display=str(case_data.get("case_name") or title),
            )
            cell.style = "Hyperlink"

    @classmethod
    def _highlight_datagram_marks(cls, sheet: Worksheet) -> None:
        """高亮HEAD/BODY分区标记单元格（黄底加粗）。"""
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value in ("HEAD", "BODY"):
                    cell.fill = cls.SHEET_STYLE_PATTERN_FILL
                    cell.font = cls.SHEET_STYLE_FONT

    # ==================== 数据源sheet样式（供数据源/性能模块视图复用） ====================

    @classmethod
    def style_data_source_sheet(cls, sheet: Worksheet) -> None:
        """
        数据源sheet统一样式：
        - 分区标记(HEAD/BODY/ASSERT_HEAD/ASSERT_BODY)所在行(垂直模式第0列标记)/列(水平模式表头行标记)整行/整列黄底；
        - 全部单元格水平/垂直居中、四边边框并自动换行，统一行高；
        - 列宽按内容自适应（[SHEET_STYLE_COL_WIDTH_MIN, SHEET_STYLE_COL_WIDTH_MAX]）；
        - 前导'协议标记转为 Excel 原生文本前缀（quotePrefix 角标）。
        """
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        if max_row < 1 or max_col < 1:
            return
        marker_rows, marker_cols = cls._locate_section_markers(sheet, max_row, max_col)
        for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            sheet.row_dimensions[row[0].row].height = cls.SHEET_STYLE_ROW_HEIGHT
            for cell in row:
                cell.alignment = cls.SHEET_STYLE_ALIGNMENT
                cell.border = cls.SHEET_STYLE_MARK_BORDER
                if cell.row in marker_rows or cell.column in marker_cols:
                    cell.fill = cls.SHEET_STYLE_PATTERN_FILL
                    cell.font = cls.SHEET_STYLE_FONT
                if isinstance(cell.value, str) and cell.value.startswith("'"):
                    cell.value = cell.value[1:]
                    cell.quotePrefix = True
        cls._auto_size_sheet_columns(sheet)

    @classmethod
    def _locate_section_markers(cls, sheet: Worksheet, max_row: int, max_col: int) -> Tuple[Set[int], Set[int]]:
        """定位垂直模式标记行（第1列）与水平模式标记列（第1行）。"""
        marker_rows = {
            row_idx for row_idx in range(2, max_row + 1)
            if str(sheet.cell(row=row_idx, column=1).value or "").strip().upper() in cls.DATAGRAM_SECTION_MARKS
        }
        marker_cols = {
            col_idx for col_idx in range(2, max_col + 1)
            if str(sheet.cell(row=1, column=col_idx).value or "").strip().upper() in cls.DATAGRAM_SECTION_MARKS
        }
        return marker_rows, marker_cols

    # ==================== 通道 B：导出接口明细（HTTP 14列 / TCP 16列双布局模板） ====================

    @classmethod
    def build_script_file_name(cls, username: Optional[str]) -> str:
        return cls._file_name(username, "接口详情数据")

    @classmethod
    async def prepare_script_export_rows(cls, case_ids: List[int], services: Any) -> Tuple[List[Dict[str, str]], List[Dict[str, Any]]]:
        loaded, invalid = await cls._load_public_api_cases(case_ids, services)
        rows: List[Dict[str, str]] = []
        for item in loaded:
            case, step = item["case"], item["step"]
            case_id, case_name, is_http = item["case_id"], item["case_name"], item["is_http"]
            args = cls._enum_val(getattr(step, "request_args_type", None))
            method = cls._enum_val(getattr(step, "request_method", None))

            problems: List[str] = []
            form_attr = CaseScriptCodec.FORM_FIELD_BY_ARGS_TYPE.get(args)
            header_text = (
                CaseScriptCodec.encode_kv_lines(getattr(step, "request_header", None), column="请求头", problems=problems)
                if is_http else ""
            )
            variables_text = CaseScriptCodec.encode_kv_lines(
                getattr(step, "defined_variables", None), column="变量", problems=problems
            )
            if form_attr:
                body_text = CaseScriptCodec.encode_kv_lines(
                    getattr(step, form_attr, None), column="请求体", problems=problems
                )
            else:
                body_text = CaseScriptCodec.encode_step_body(step)
            problems.extend(CaseScriptCodec.check_extract_roundtrip(getattr(step, "extract_variables", None)))
            problems.extend(CaseScriptCodec.check_assert_roundtrip(getattr(step, "assert_validators", None)))

            # 往返不安全项直接拦截该用例，避免导出成功但导入错位
            if problems:
                invalid.append({"case_id": case_id, "case_name": case_name, "reason": "; ".join(problems)})
                continue

            row = {
                "接口名称": case_name,
                "所属应用": item["project_name"],
                "协议类型": cls.HTTP_PROTOCOL if is_http else cls.TCP_PROTOCOL,
                "接口描述": getattr(case, "case_desc", None) or "",
                "请求方式": method if is_http else "",
                "配置名称": getattr(step, "request_config_name", None) or "",
                "请求路径": (getattr(step, "request_url", None) or "") if is_http else "",
                "请求体类型": args,
                "请求体": body_text,
                "请求头": header_text,
                "变量": variables_text,
                "提取": CaseScriptCodec.encode_extract_lines(getattr(step, "extract_variables", None)),
                "断言": CaseScriptCodec.encode_assert_lines(getattr(step, "assert_validators", None)),
                "所属人员": getattr(case, "owner_user", None) or getattr(case, "created_user", None) or "",
            }
            if not is_http:
                row.update(cls._tcp_frame_export_values(step))
            rows.append(row)
        LOGGER.info(f"导出脚本准备完成: 有效{len(rows)}个, 无效{len(invalid)}个")
        return rows, invalid

    @classmethod
    def _tcp_frame_export_values(cls, step: Any) -> Dict[str, str]:
        """
        TCP帧配置导出五列。

        仅导出执行引擎成帧必需的五字段；连接方式/编码/HEX开关/响应超时不进模板，导入更新时保留库内原值。
        """
        return {
            column: str(cls._enum_val(getattr(step, attr, None)))
            for column, attr in cls.TCP_FRAME_CONFIG_COLUMNS
        }

    @staticmethod
    def _script_template_path(file_name: str) -> str:
        return os.path.join(PROJECT_CONFIG.OUTPUT_DIR, "template", file_name)

    @classmethod
    def build_script_workbook(cls, rows: List[Dict[str, str]]) -> Workbook:
        """
        按导出用例的协议组合选择模板：全HTTP用HTTP模板、全TCP用TCP模板、混合用HTTP&TCP模板；
        模板内HTTP行写入[HTTP模板]sheet、TCP行写入[TCP模板]sheet；填充完成后将用户可见页签
        重命名为[HTTP数据]/[TCP数据]，[填写说明]sheet不做任何修改。

        :param rows: 行字典列表（键为列名，TCP行额外携带帧配置列）
        :return: 填充完成的工作簿
        :raises RuntimeError: 模板缺失或模板表头/页签与预定义布局不一致时
        """
        has_http = any(row.get("协议类型") == cls.HTTP_PROTOCOL for row in rows)
        has_tcp = any(row.get("协议类型") == cls.TCP_PROTOCOL for row in rows)
        template_name = cls.SCRIPT_TEMPLATE_FILES.get((has_http, has_tcp))
        if template_name is None:
            # 空行列表由上层拦截（无可导出的合规用例），此处仅防御
            raise RuntimeError("无可导出的合规用例")

        template_path = cls._script_template_path(template_name)
        if not os.path.isfile(template_path):
            LOGGER.error(f"脚本模板文件不存在: {template_path}")
            raise RuntimeError("脚本模板文件不存在, 请联系管理员")

        workbook = load_workbook(template_path)
        for protocol, (sheet_name, columns) in cls.SCRIPT_SHEET_LAYOUTS.items():
            protocol_rows = [row for row in rows if row.get("协议类型") == protocol]
            if not protocol_rows:
                continue
            sheet = cls._resolve_script_sheet(workbook, template_name, sheet_name, columns)
            cls._fill_script_sheet(sheet, columns, protocol_rows)
            sheet.title = cls.SCRIPT_EXPORT_SHEET_NAMES[protocol]
        return workbook

    @staticmethod
    def _resolve_script_sheet(
            workbook: Workbook, template_name: str, sheet_name: str, columns: Tuple[str, ...]
    ) -> Worksheet:
        """校验模板页签存在且表头与布局注册表一致，返回该sheet。"""
        if sheet_name not in workbook.sheetnames:
            LOGGER.error(f"脚本模板缺少页签[{sheet_name}]: 模板={template_name}, 页签={workbook.sheetnames}")
            raise RuntimeError(f"脚本模板缺少页签[{sheet_name}], 请联系管理员")
        sheet = workbook[sheet_name]
        header = [cell.value for cell in sheet[1][: len(columns)]]
        if header != list(columns):
            LOGGER.error(f"脚本模板[{sheet_name}]表头与预定义表头不一致: 模板={header}, 预定义={list(columns)}")
            raise RuntimeError(f"脚本模板[{sheet_name}]表头与预定义表头不一致, 请联系管理员")
        return sheet

    @classmethod
    def _fill_script_sheet(cls, sheet: Worksheet, columns: Tuple[str, ...], rows: List[Dict[str, str]]) -> None:
        """按列定义写入数据行，继承模板数据行字体并套用数据区样式与列宽。"""
        # 新增行不会继承模板字体，需以首条数据行为基准，避免回退到工作簿默认字号。
        data_row_fonts = [
            copy(sheet.cell(row=cls.SCRIPT_DATA_START_ROW, column=col_index).font)
            for col_index in range(1, len(columns) + 1)
        ]
        for row_index, row in enumerate(rows, start=cls.SCRIPT_DATA_START_ROW):
            for col_index, (column, data_font) in enumerate(zip(columns, data_row_fonts), start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=row.get(column) or "")
                # 工作簿按字体值登记并去重，各单元格共享字体即可，无需逐格复制。
                cell.font = data_font
        cls._set_sheet_cells_style(sheet, start_row=cls.SCRIPT_DATA_START_ROW, row_height=cls.SHEET_STYLE_ROW_HEIGHT)
        cls._auto_size_sheet_columns(sheet)

    # ==================== 通道 C：导入接口明细（解析/校验/落库） ====================

    @staticmethod
    def _cell_text(value: Any) -> str:
        """单元格值转文本：整数值浮点去掉小数尾巴，其余去首尾空白。"""
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def _detect_script_sheet_layout(cls, sheet: Worksheet) -> Optional[Tuple[str, Tuple[str, ...]]]:
        """
        按首个sheet表头识别模板协议布局，兼容尾随空列；识别失败返回None。

        :param sheet: 工作簿第一个sheet
        :return: (协议, 列定义) 或 None
        """
        max_known = max(len(columns) for _, columns in cls.SCRIPT_SHEET_LAYOUTS.values())
        header = [cls._cell_text(cell.value) for cell in sheet[1][:max_known]]
        while header and not header[-1]:
            header.pop()
        for protocol, (_, columns) in cls.SCRIPT_SHEET_LAYOUTS.items():
            if header == list(columns):
                return protocol, columns
        return None

    @classmethod
    def _parse_tcp_frame_fields(cls, cells: Dict[str, str]) -> Tuple[Dict[str, Any], List[str]]:
        """
        解析并校验TCP帧配置五列（报文进制/长度位/报文总长度/请求结束符/响应结束符）。

        校验口径与执行引擎 resolve_tcp_frame_config 对齐：报文进制≠不使用时长度位与报文总长度必填；
        报文进制取值与 AutoTestTcpRadix 枚举一致（不使用/10进制/16进制），不再做别名映射。

        :param cells: 行单元格字典（键为列名）
        :return: (可落库的TCP帧字段字典, 错误列表)
        """
        errors: List[str] = []
        radix_raw = cells["报文进制"]
        radix = radix_raw or None
        if not radix:
            errors.append("TCP协议[报文进制]不允许为空")
        elif radix not in cls.TCP_RADIX_VALUES:
            errors.append(f"报文进制 {radix_raw} 不被允许")
            radix = None

        radix_length: Optional[int] = None
        radix_length_raw = cells["长度位"]
        if radix_length_raw:
            if not re.fullmatch(r"\d+", radix_length_raw):
                errors.append(f"长度位 {radix_length_raw} 必须为整数")
            elif not (cls.TCP_RADIX_LENGTH_MIN <= int(radix_length_raw) <= cls.TCP_RADIX_LENGTH_MAX):
                errors.append(f"长度位 {radix_length_raw} 超出允许范围[{cls.TCP_RADIX_LENGTH_MIN}, {cls.TCP_RADIX_LENGTH_MAX}]")
            else:
                radix_length = int(radix_length_raw)
        use_radix = radix is not None and radix != AutoTestTcpRadix.NONE.value
        if use_radix and not radix_length_raw:
            errors.append(f"报文进制为[{radix}]时, 长度位不允许为空")

        datagram_length = cells["报文总长度"] or None
        if datagram_length and datagram_length not in cls.TCP_DATAGRAM_LENGTH_VALUES:
            errors.append(f"报文总长度 {datagram_length} 不被允许")
            datagram_length = None
        if use_radix and not datagram_length:
            errors.append(f"报文进制为[{radix}]时, 报文总长度不允许为空")

        request_end_symbol = cells["请求结束符"] or None
        if request_end_symbol and request_end_symbol not in cls.TCP_REQUEST_END_SYMBOL_VALUES:
            errors.append(f"请求结束符 {request_end_symbol} 不被允许")
            request_end_symbol = None

        response_end_symbol = cells["响应结束符"]
        if not response_end_symbol or response_end_symbol == AutoTestTcpEndSymbol.NONE.value:
            response_end_symbol = None
        elif len(response_end_symbol) > cls.TCP_RESPONSE_END_SYMBOL_MAX:
            errors.append(f"响应结束符长度超过{cls.TCP_RESPONSE_END_SYMBOL_MAX}个字符")
            response_end_symbol = None

        return {
            "tcp_radix": radix,
            "tcp_radix_length": radix_length,
            "tcp_datagram_length": datagram_length,
            "tcp_request_end_symbol": request_end_symbol,
            "tcp_response_end_symbol": response_end_symbol,
        }, errors

    @classmethod
    def _validate_row_identity(
            cls, cells: Dict[str, str], sheet_protocol: str
    ) -> Tuple[str, str, Optional[str], List[str]]:
        """校验接口名称/所属应用必填与协议类型合法性，返回(名称, 应用, 协议或None, 错误)。"""
        errors: List[str] = []
        case_name, project_name = cells["接口名称"], cells["所属应用"]
        protocol_raw = cells["协议类型"]
        protocol: Optional[str] = protocol_raw.upper()
        if not case_name:
            errors.append("接口名称不允许为空")
        if not project_name:
            errors.append("所属应用不允许为空")
        if protocol not in (cls.HTTP_PROTOCOL, cls.TCP_PROTOCOL):
            errors.append(f"协议类型 {protocol_raw} 不被允许")
            protocol = None
        elif protocol != sheet_protocol:
            errors.append(f"协议类型 {protocol} 与模板页签不一致, 当前页签仅支持{sheet_protocol}协议")
            protocol = None
        return case_name, project_name, protocol, errors

    @classmethod
    def _validate_http_columns(cls, cells: Dict[str, str]) -> Tuple[str, str, str, List[str]]:
        """校验HTTP协议专属列，返回(请求方式, 请求路径, 请求头文本, 错误)。"""
        errors: List[str] = []
        method, request_url, header_text = cells["请求方式"].upper(), cells["请求路径"], cells["请求头"]
        if not method:
            errors.append("HTTP协议[请求方式]不允许为空")
        elif method not in cls.HTTP_METHODS:
            errors.append(f"请求方式 {method} 不被允许")
        if not request_url:
            errors.append("HTTP协议[请求路径]不允许为空")
        return method, request_url, header_text, errors

    @classmethod
    def _validate_args_type(cls, cells: Dict[str, str], sheet_protocol: str) -> Tuple[str, List[str]]:
        """校验请求体类型：全局合法且TCP仅支持json/xml/raw。"""
        errors: List[str] = []
        args_type = cells["请求体类型"]
        if not args_type:
            errors.append("请求体类型不允许为空")
        elif args_type not in cls.REQ_ARGS_TYPE:
            errors.append(f"请求体类型 {args_type} 不被允许")
        elif sheet_protocol == cls.TCP_PROTOCOL and args_type not in cls.TCP_ARGS_TYPE:
            errors.append(f"TCP协议[请求体类型]仅支持 {'/'.join(sorted(cls.TCP_ARGS_TYPE))}")
        return args_type, errors

    @classmethod
    def _parse_script_row(cls, cells: Dict[str, str], row_no: int, sheet_protocol: str) -> Tuple[Optional[Dict[str, Any]], List[str]]:
        """
        解析单行数据：先校验协议与页签一致性，再按布局分支校验协议专属列，最后解析报文列。

        错误按"身份列 → 协议专属列 → 请求体 → 报文列"顺序收集，与模板列顺序一致，便于用户逐列修正。

        :param cells: 行单元格字典（键为列名）
        :param row_no: 行号（仅用于错误定位信息，由调用方拼装）
        :param sheet_protocol: 首个sheet识别出的协议布局
        :return: (行字典, 错误列表)；存在错误时行字典为None
        """
        case_name, project_name, protocol, errors = cls._validate_row_identity(cells, sheet_protocol)

        if sheet_protocol == cls.HTTP_PROTOCOL:
            method, request_url, header_text, http_errors = cls._validate_http_columns(cells)
            errors.extend(http_errors)
            tcp_fields: Dict[str, Any] = {attr: None for _, attr in cls.TCP_FRAME_CONFIG_COLUMNS}
        else:
            method, request_url, header_text = "", "", ""
            tcp_fields, tcp_errors = cls._parse_tcp_frame_fields(cells)
            errors.extend(tcp_errors)

        args_type, args_errors = cls._validate_args_type(cells, sheet_protocol)
        errors.extend(args_errors)

        args_ok = (
                protocol is not None
                and args_type in cls.REQ_ARGS_TYPE
                and (sheet_protocol == cls.HTTP_PROTOCOL or args_type in cls.TCP_ARGS_TYPE)
        )
        body_fields: Dict[str, Any] = {field: None for field in CaseScriptCodec.BODY_STEP_FIELDS}
        if args_ok:
            body_fields, body_errors = CaseScriptCodec.decode_body(args_type, cells["请求体"])
            errors.extend(body_errors)

        request_header: Optional[List[Dict[str, str]]] = None
        if sheet_protocol == cls.HTTP_PROTOCOL:
            request_header, header_errors = CaseScriptCodec.decode_kv_lines(header_text, "请求头")
            errors.extend(header_errors)
            request_header = request_header or None
        defined_variables, kv_errors = CaseScriptCodec.decode_kv_lines(cells["变量"], "变量")
        errors.extend(kv_errors)
        extract_variables, extract_errors = CaseScriptCodec.decode_extract_lines(cells["提取"])
        errors.extend(extract_errors)
        assert_validators, assert_errors = CaseScriptCodec.decode_assert_lines(cells["断言"])
        errors.extend(assert_errors)
        if errors:
            return None, errors
        return {
            "row_no": row_no,
            "case_name": case_name,
            "project_name": project_name,
            "protocol": protocol,
            "case_desc": cells["接口描述"] or None,
            "step_name": case_name,
            "request_method": method or None,
            "request_config_name": cells["配置名称"] or None,
            "request_url": request_url or None,
            "request_args_type": args_type,
            "request_header": request_header,
            "defined_variables": defined_variables or None,
            "extract_variables": extract_variables or None,
            "assert_validators": assert_validators or None,
            **tcp_fields,
            **body_fields,
        }, []

    @classmethod
    def parse_script_workbook(cls, content: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """
        解析导入模板首个sheet：识别布局、逐行校验，输出可落库行与不合规明细。

        仅读取首个sheet（模板填写说明页明确约定），空行跳过。

        :param content: xlsx文件字节
        :return: (行字典列表, 不合规明细[{row, reason}])
        """
        try:
            workbook = load_workbook(io.BytesIO(content), data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as e:
            # 第三方解析器对损坏/伪装文件的异常类型不统一，统一转为可读的校验结果，避免接口500
            LOGGER.error(f"解析接口模板工作簿失败: {e}\n{traceback.format_exc()}")
            return [], [{"row": 1, "reason": "模板文件无法解析, 请下载最新接口模板后重试"}]
        if not workbook.sheetnames:
            return [], [{"row": 1, "reason": "模板文件缺少工作表"}]
        sheet = workbook[workbook.sheetnames[0]]
        layout = cls._detect_script_sheet_layout(sheet)
        if layout is None:
            return [], [{"row": 1, "reason": "脚本模板表头与预定义表头不一致"}]
        sheet_protocol, columns = layout

        rows: List[Dict[str, Any]] = []
        invalid: List[Dict[str, Any]] = []
        for row_no, excel_row in enumerate(
                sheet.iter_rows(min_row=cls.SCRIPT_DATA_START_ROW, max_col=len(columns), values_only=True),
                start=cls.SCRIPT_DATA_START_ROW,
        ):
            cells = {column: cls._cell_text(value) for column, value in zip(columns, excel_row)}
            if not any(cells.values()):
                continue
            row, errors = cls._parse_script_row(cells, row_no, sheet_protocol)
            if errors:
                invalid.append({"row": row_no, "reason": "；".join(errors)})
                continue
            rows.append(row)
        if not rows and not invalid:
            invalid.append({"row": cls.SCRIPT_DATA_START_ROW, "reason": "文件无有效数据行"})
        LOGGER.info(f"导入脚本解析完成: 协议={sheet_protocol}, 有效{len(rows)}行, 无效{len(invalid)}行")
        return rows, invalid

    @classmethod
    async def parse_script_workbook_async(cls, content: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """
        异步解析入口：openpyxl 解析为阻塞IO，放入线程执行避免阻塞事件循环。

        :param content: xlsx文件字节
        :return: 同 parse_script_workbook
        """
        return await asyncio.to_thread(cls.parse_script_workbook, content)

    @classmethod
    def _build_import_step_payload(cls, item: Dict[str, Any]) -> Dict[str, Any]:
        """
        组装步骤落库载荷。

        HTTP行TCP帧五字段均为None：覆盖更新时顺带清空存量TCP帧配置，避免协议切换后残留脏配置；
        连接方式/编码/HEX开关/响应超时不在模板内，这里不触碰以保留库内原值。

        :param item: 解析并匹配完成的行字典
        :return: 步骤创建/更新载荷
        """
        return {
            "step_name": item["step_name"],
            "step_desc": item["case_desc"] or "",
            "step_type": AutoTestStepType.HTTP if item["protocol"] == cls.HTTP_PROTOCOL else AutoTestStepType.TCP,
            "request_project_id": item["project_id"],
            "request_method": item["request_method"],
            "request_url": item["request_url"],
            "request_config_name": item["request_config_name"],
            "request_args_type": AutoTestReqArgsType(item["request_args_type"]),
            "request_body": item["request_body"],
            "request_text": item["request_text"],
            "request_header": item["request_header"],
            "request_params": item["request_params"],
            "request_form_data": item["request_form_data"],
            "request_form_urlencoded": item["request_form_urlencoded"],
            "defined_variables": item["defined_variables"],
            "extract_variables": item["extract_variables"],
            "assert_validators": item["assert_validators"],
            "tcp_radix": AutoTestTcpRadix(item["tcp_radix"]) if item.get("tcp_radix") else None,
            "tcp_radix_length": item.get("tcp_radix_length"),
            "tcp_datagram_length": (
                AutoTestTcpDatagramLength(item["tcp_datagram_length"]) if item.get("tcp_datagram_length") else None
            ),
            "tcp_request_end_symbol": (
                AutoTestTcpEndSymbol(item["tcp_request_end_symbol"]) if item.get("tcp_request_end_symbol") else None
            ),
            "tcp_response_end_symbol": item.get("tcp_response_end_symbol") or None,
        }

    @staticmethod
    def _find_duplicate_rows(rows: List[Dict[str, Any]]) -> Tuple[Set[int], List[Dict[str, Any]]]:
        """同应用同接口名称只保留首行，返回重复行号集合与不合规明细。"""
        seen: Dict[Tuple[str, str], int] = {}
        duplicate_row_nos: Set[int] = set()
        invalid: List[Dict[str, Any]] = []
        for row in rows:
            key = (row["project_name"], row["case_name"])
            if key in seen:
                invalid.append({
                    "row": row["row_no"],
                    "reason": f"与第{seen[key]}行重复, 同应用同名称只允许保留一条",
                })
                duplicate_row_nos.add(row["row_no"])
            else:
                seen[key] = row["row_no"]
        return duplicate_row_nos, invalid

    @staticmethod
    async def _load_import_project_ids(
            rows: List[Dict[str, Any]], services: Any
    ) -> Dict[str, int]:
        """按模板出现的应用名批量定位启用应用，返回 应用名→主键（缺失即未命中）。"""
        project_names = list(dict.fromkeys(row["project_name"] for row in rows))
        if not project_names:
            return {}
        projects: List[AutoTestProjectModel] = await services.project_curd.model.filter(
            project_name__in=project_names, state__not=1
        ).all()
        project_ids: Dict[str, int] = {}
        for project in projects:
            project_ids.setdefault(project.project_name, project.id)
        return project_ids

    @staticmethod
    async def _load_import_cases(
            rows: List[Dict[str, Any]],
            project_ids: Dict[str, int],
            owner_user: Optional[str],
    ) -> Dict[Tuple[int, str], List[AutoTestCaseModel]]:
        """
        批量定位同应用+同名称+同所属人的公共接口（含软删除，用于恢复覆盖）。

        :return: (应用ID, 接口名称) → 同名接口列表
        """
        grouped: Dict[Tuple[int, str], List[AutoTestCaseModel]] = defaultdict(list)
        if not owner_user:
            return grouped
        case_project_ids = sorted(set(project_ids.values()))
        case_names = list(dict.fromkeys(row["case_name"] for row in rows))
        if not case_project_ids or not case_names:
            return grouped
        cases: List[AutoTestCaseModel] = await AutoTestCaseModel.filter(
            case_project__in=case_project_ids,
            case_name__in=case_names,
            case_type=AutoTestCaseType.PUBLIC_API,
            owner_user=owner_user,
        ).all()
        for case in cases:
            grouped[(case.case_project, case.case_name)].append(case)
        return grouped

    @staticmethod
    async def _load_import_root_steps(
            grouped_cases: Dict[Tuple[int, str], List[AutoTestCaseModel]],
    ) -> Dict[int, List[AutoTestStepModel]]:
        """
        批量装载命中接口的根步骤（含软删，供软删接口恢复覆盖时计数）。

        :return: 用例ID → 根步骤列表
        """
        case_ids = [case.id for cases in grouped_cases.values() for case in cases]
        if not case_ids:
            return {}
        steps: List[AutoTestStepModel] = await AutoTestStepModel.filter(
            case_id__in=case_ids, parent_step_id=None
        ).all()
        grouped: Dict[int, List[AutoTestStepModel]] = defaultdict(list)
        for step in steps:
            grouped[step.case_id].append(step)
        return grouped

    @classmethod
    async def _match_import_rows(
            cls,
            rows: List[Dict[str, Any]],
            services: Any,
            owner_user: Optional[str],
    ) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """
        批量行级匹配：应用、同名接口、根步骤各一次查询，再按行判定合规模板。

        匹配口径：所属应用启用 + 同应用/同名称/同所属人的公共接口（含软删）；命中接口须恰有一个根步骤。

        :param rows: 去重后的解析行列表
        :param services: CRUD聚合
        :param owner_user: 当前登录所属人账号
        :return: (匹配结果列表, 不合规明细[{row, reason}])
        """
        project_ids = await cls._load_import_project_ids(rows, services)
        grouped_cases = await cls._load_import_cases(rows, project_ids, owner_user)
        root_steps_by_case = await cls._load_import_root_steps(grouped_cases)

        prepared: List[Dict[str, Any]] = []
        invalid: List[Dict[str, Any]] = []
        for row in rows:
            errors: List[str] = []
            project_id = project_ids.get(row["project_name"])
            if project_id is None:
                errors.append(f"所属应用 {row['project_name']} 不存在")
            existing_case: Optional[AutoTestCaseModel] = None
            existing_step: Optional[AutoTestStepModel] = None
            if project_id is not None:
                if not owner_user:
                    errors.append("当前登录账号为空, 无法获取登录用户并定位公共接口")
                else:
                    matched_cases = grouped_cases.get((project_id, row["case_name"]), [])
                    if len(matched_cases) > 1:
                        errors.append(f"应用[{row['project_name']}]存在多条同名接口, 请先处理同名接口: {row['case_name']}")
                    elif matched_cases:
                        existing_case = matched_cases[0]
                        root_steps = root_steps_by_case.get(existing_case.id, [])
                        # 软删接口恢复覆盖时需连同软删步骤一起定位，否则根步骤计数失真
                        if existing_case.state != 1:
                            root_steps = [step for step in root_steps if step.state != 1]
                        if len(root_steps) != 1:
                            errors.append(f"接口[{row['case_name']}]步骤形态异常, 根步骤数为{len(root_steps)}")
                        else:
                            existing_step = root_steps[0]
            if errors:
                invalid.append({"row": row["row_no"], "reason": "；".join(errors)})
                continue
            prepared.append({
                **row,
                "project_id": project_id,
                "existing_case": existing_case,
                "existing_step": existing_step,
            })
        return prepared, invalid

    @classmethod
    async def _persist_import_rows(
            cls,
            prepared: List[Dict[str, Any]],
            services: Any,
            owner_user: Optional[str],
    ) -> Tuple[int, int]:
        """
        单事务落库：不存在则新增用例+步骤，存在则更新或恢复覆盖（版本号自增）。

        :return: (新增数, 更新数)
        """
        created = updated = 0
        async with in_transaction():
            for item in prepared:
                step_payload = cls._build_import_step_payload(item)
                if item["existing_case"] is None:
                    new_case = await services.case_curd.create(obj_in={
                        "case_name": item["case_name"],
                        "case_project": item["project_id"],
                        "case_type": AutoTestCaseType.PUBLIC_API,
                        "case_attr": AutoTestCaseAttr.TRUE_CASE,
                        "case_desc": item["case_desc"],
                        "case_tags": None,
                        "session_variables": None,
                        "case_steps": 1,
                        "owner_user": owner_user,
                    })
                    await services.step_curd.create(obj_in={"case_id": new_case.id, "step_no": 1, **step_payload})
                    created += 1
                else:
                    existing_case = item["existing_case"]
                    await services.case_curd.update(
                        id=existing_case.id,
                        obj_in={
                            "case_desc": item["case_desc"],
                            "case_tags": None,
                            "case_version": (getattr(existing_case, "case_version", None) or 1) + 1,
                            "state": 0,
                        },
                    )
                    await services.step_curd.update(
                        id=item["existing_step"].id,
                        obj_in={**step_payload, "state": 0},
                    )
                    updated += 1
        return created, updated

    @classmethod
    async def import_script_rows(cls, rows: List[Dict[str, Any]], services: Any) -> Tuple[Dict[str, int], List[Dict[str, Any]]]:
        """
        按所属应用、接口名称、当前登录所属人匹配公共接口，含软删，存在则更新或恢复覆盖、不存在则新增，校验通过后单事务落库。

        全有或全无：任一行匹配不合规即整批取消落库，返回不合规明细供用户修正后重试。

        :param rows: 解析通过的行列表
        :param services: CRUD聚合
        :return: ({"created_count": 新增数, "updated_count": 更新数}, 不合规明细)
        """
        owner_user = get_current_username()
        duplicate_row_nos, invalid = cls._find_duplicate_rows(rows)
        pending_rows = [row for row in rows if row["row_no"] not in duplicate_row_nos]
        prepared, match_invalid = await cls._match_import_rows(pending_rows, services, owner_user)
        invalid.extend(match_invalid)
        if invalid:
            return {}, invalid

        created, updated = await cls._persist_import_rows(prepared, services, owner_user)
        LOGGER.info(f"导入脚本落库完成: 新增{created}个, 更新{updated}个")
        return {"created_count": created, "updated_count": updated}, []
